import json
from datetime import datetime
from unittest.mock import patch

import yaml
from click.testing import CliRunner
from openpyxl import load_workbook

from filtertool.analysis import build_crosstab
from filtertool.classify.classifier import classify_papers
from filtertool.cli import cli
from filtertool.coding import auto_manual_agreement, import_manual_coding
from filtertool.export import export_crosstab_to_excel
from filtertool.filter.rule_filter import apply_rule_filter
from filtertool.models import Paper, PaperStatus
from filtertool.pipeline import Pipeline
from filtertool.prisma import build_prisma_counts, select_search_events
from filtertool.provenance import append_search_log, config_hash
from filtertool.rate_limit import AdaptiveRateLimiter, get_rate_limiter
from filtertool.screening import apply_screening_rows, export_screening_sheet
from filtertool.search.base import BaseSearchAdapter
from filtertool.search.crossref import CrossrefAdapter
from filtertool.search.semantic_scholar import SemanticScholarAdapter
from filtertool.storage import PaperStore
from filtertool.verify import _make_request

CONFIG = {
    "rule_filter": {
        "min_keyword_hits": 2,
        "llm_terms": ["large language model", "LLM", "AI agent"],
        "offensive_terms": ["penetration test", "exploit", "CTF"],
        "exclusion_keywords": ["protein folding"],
    },
    "taxonomy": {
        "attack_methods": {
            "penetration_testing": {"keywords": ["penetration test", "pentest"]},
            "ctf_solving": {"keywords": ["capture the flag", "ctf"]},
        },
        "attack_stages": {"exploitation": {"keywords": ["exploit"]}},
        "optimization_techniques": {"tool_use": {"keywords": ["tool use", "nmap"]}},
        "optimization_objectives": {"success_rate": {"keywords": ["success rate"]}},
    },
    "analysis": {"sparse_max_k": 1, "min_marginal_total": 1, "min_total_papers": 2},
}


def test_rule_filter_requires_llm_and_offensive_terms():
    llm_only = Paper(title="Large language model benchmark", status=PaperStatus.DEDUPED.value)
    both_groups = Paper(
        title="LLM agent for autonomous penetration testing",
        status=PaperStatus.DEDUPED.value,
    )
    noisy = Paper(
        title="Large language model for protein folding",
        status=PaperStatus.DEDUPED.value,
    )

    apply_rule_filter([llm_only, both_groups, noisy], CONFIG)

    assert llm_only.status == PaperStatus.REVIEW_NEEDED.value
    assert both_groups.status == PaperStatus.RULE_INCLUDED.value
    assert noisy.status == PaperStatus.REVIEW_NEEDED.value
    assert "protein folding" in noisy.screening_decisions[-1]["reason"]


def test_rule_filter_rescreens_automated_results_but_preserves_human_decisions():
    candidate = Paper(
        title="Large language model for penetration testing",
        status=PaperStatus.SEMANTIC_LOW.value,
    )
    unrelated = Paper(
        title="A study of clinical imaging",
        status=PaperStatus.SEMANTIC_HIGH.value,
    )
    reviewed = Paper(
        title="Human-reviewed paper",
        status=PaperStatus.HUMAN_EXCLUDED.value,
        human_decision="Exclude",
        reason_code="EC2",
    )
    reviewed.add_screening_decision(
        "human_screening:title_abstract", "exclude", "Out of scope", reason_code="EC2"
    )

    apply_rule_filter([candidate, unrelated, reviewed], CONFIG)

    assert candidate.status == PaperStatus.RULE_INCLUDED.value
    assert unrelated.status == PaperStatus.REVIEW_NEEDED.value
    assert reviewed.status == PaperStatus.HUMAN_EXCLUDED.value
    assert reviewed.human_decision == "Exclude"


def test_semantic_filter_rechecks_existing_semantic_statuses(monkeypatch):
    from filtertool.filter import semantic_filter

    class Scores:
        def __init__(self, values):
            self._values = values

        def max(self, dim):
            return self

        @property
        def values(self):
            return self

        def tolist(self):
            return self._values

    class Model:
        def encode(self, texts, convert_to_tensor=True):
            return texts

    class Similarity:
        @staticmethod
        def cos_sim(papers, references):
            return Scores([0.9 for _ in papers])

    paper = Paper(
        title="LLM agent automates penetration testing",
        status=PaperStatus.SEMANTIC_LOW.value,
    )
    paper.keyword_hit_count = 2
    monkeypatch.setattr(semantic_filter, "SentenceTransformer", Model)
    monkeypatch.setattr(semantic_filter, "_get_model", lambda _: Model())
    monkeypatch.setattr(semantic_filter, "util", Similarity)

    semantic_filter.apply_semantic_filter(
        [paper],
        {
            "semantic_filter": {
                "reference_sentences": ["LLM agents for penetration testing"],
                "high_threshold": 0.7,
                "low_threshold": 0.3,
                "semantic_weight": 0.8,
                "keyword_weight": 0.2,
            },
            "rule_filter": {"min_keyword_hits": 2},
        },
    )

    assert paper.semantic_label == "HIGH"
    assert paper.status == PaperStatus.SEMANTIC_HIGH.value


def test_screening_import_validates_batch_before_mutation():
    paper = Paper(title="Candidate")
    rows = [
        {"id": paper.id, "human_decision": "Exclude", "reason_code": "EC1"},
        {"id": "unknown", "human_decision": "Include", "reason_code": "IC1"},
    ]

    try:
        apply_screening_rows([paper], rows)
    except ValueError:
        pass
    else:
        raise AssertionError("unknown paper id should reject the import")

    assert paper.status == PaperStatus.NEW.value
    assert paper.screening_decisions == []


def test_human_decisions_drive_status_and_preserve_reason_codes():
    paper = Paper(title="Candidate")
    apply_screening_rows(
        [paper],
        [{"id": paper.id, "human_decision": "Include", "reason_code": "IC2"}],
    )
    assert paper.status == PaperStatus.HUMAN_INCLUDED.value
    assert paper.screening_decisions[-1]["reason_code"] == "IC2"

    apply_screening_rows(
        [paper],
        [{
            "id": paper.id,
            "screening_stage": "full_text",
            "human_decision": "Exclude",
            "reason_code": "EC2",
            "fulltext_retrieved": "yes",
        }],
    )
    assert paper.status == PaperStatus.FULLTEXT_EXCLUDED.value
    assert paper.screening_decisions[-1]["reason_code"] == "EC2"


def test_screening_import_skips_rows_without_decisions():
    reviewed = Paper(title="Reviewed candidate")
    pending = Paper(title="Not reviewed yet")

    imported = apply_screening_rows(
        [reviewed, pending],
        [
            {"id": reviewed.id, "human_decision": "Exclude", "reason_code": "EC1"},
            {"id": pending.id, "human_decision": ""},
        ],
    )

    assert imported == 1
    assert reviewed.status == PaperStatus.HUMAN_EXCLUDED.value
    assert pending.status == PaperStatus.NEW.value


def test_title_unsure_reappears_until_resolved(tmp_path):
    paper = Paper(title="Ambiguous LLM candidate", status=PaperStatus.REVIEW_NEEDED.value)
    apply_screening_rows(
        [paper],
        [{"id": paper.id, "screening_stage": "title_abstract", "human_decision": "Unsure"}],
    )
    path = tmp_path / "screening.xlsx"
    export_screening_sheet([paper], path, stage="title_abstract")
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["Screening"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    assert rows[0][0] == paper.id

    apply_screening_rows(
        [paper],
        [{"id": paper.id, "screening_stage": "title_abstract", "human_decision": "Include", "reason_code": "IC1"}],
    )
    export_screening_sheet([paper], path, stage="title_abstract")
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook["Screening"].max_row == 1
    workbook.close()
    assert paper.status == PaperStatus.HUMAN_INCLUDED.value


def test_screening_sheet_has_fixed_dependent_reason_dropdowns(tmp_path):
    path = tmp_path / "screening.xlsx"
    export_screening_sheet([Paper(title="Candidate")], path)

    workbook = load_workbook(path)
    sheet = workbook["Screening"]
    validations = list(sheet.data_validations.dataValidation)

    assert any(validation.formula1 == '"Include,Exclude,Unsure"' for validation in validations)
    reason_validation = next(
        validation for validation in validations
        if validation.type == "list" and "N2" in str(validation.sqref)
    )
    assert "IncludeReasonCodes" in reason_validation.formula1
    assert "ExcludeReasonCodes" in reason_validation.formula1
    codebook = workbook["Codebook"]
    assert codebook["A2"].value == "IC1"
    assert codebook["A3"].value == "IC2"
    assert codebook["B2"].value == "EC1"
    assert codebook["B4"].value == "EC3"
    workbook.close()


def test_screening_date_updates_on_decision_or_reason_change_only():
    paper = Paper(title="Candidate")
    apply_screening_rows(
        [paper], [{"id": paper.id, "human_decision": "Include", "reason_code": "IC1"}]
    )
    assert paper.review_date == datetime.now().astimezone().date().isoformat()

    paper.review_date = "2000-01-01"
    apply_screening_rows(
        [paper], [{"id": paper.id, "human_decision": "Include", "reason_code": "IC1"}]
    )
    assert paper.review_date == "2000-01-01"

    apply_screening_rows(
        [paper], [{"id": paper.id, "human_decision": "Exclude", "reason_code": "EC1"}]
    )
    assert paper.review_date == datetime.now().astimezone().date().isoformat()


def test_screening_import_rejects_reason_codes_outside_codebook():
    paper = Paper(title="Candidate")
    try:
        apply_screening_rows(
            [paper], [{"id": paper.id, "human_decision": "Exclude", "reason_code": "EC9"}]
        )
    except ValueError as error:
        assert "reason_code" in str(error)
    else:
        raise AssertionError("reason codes outside the fixed codebook should be rejected")

    assert paper.status == PaperStatus.NEW.value


def test_prisma_counts_unresolved_unsure_separately():
    title_unsure = Paper(status=PaperStatus.REVIEW_NEEDED.value)
    title_unsure.add_screening_decision("human_screening:title_abstract", "unsure", "review")
    full_text_unsure = Paper(status=PaperStatus.FULLTEXT_ASSESSED.value)
    full_text_unsure.add_screening_decision("human_screening:title_abstract", "include", "IC1")
    full_text_unsure.add_screening_decision("human_screening:full_text", "unsure", "review")

    counts = build_prisma_counts(
        [title_unsure, full_text_unsure],
        search_events=[{"source": "openalex", "n_returned": 2}],
    )

    assert counts["title_abstract_awaiting_resolution"] == 1
    assert counts["full_text_unsure"] == 1
    assert counts["full_text_awaiting_resolution"] == 1
    assert counts["counts_are_consistent"] is False


def test_classifier_only_suggests_for_human_included_papers():
    automatic_candidate = Paper(
        title="LLM agent uses nmap to exploit systems",
        status=PaperStatus.SEMANTIC_HIGH.value,
    )
    human_candidate = Paper(
        title="LLM agent uses nmap to exploit systems",
        status=PaperStatus.HUMAN_INCLUDED.value,
    )

    classify_papers([automatic_candidate, human_candidate], CONFIG)

    assert automatic_candidate.attack_methods_auto == []
    assert human_candidate.status == PaperStatus.HUMAN_INCLUDED.value
    assert human_candidate.optimization_techniques_auto == ["tool_use"]
    assert human_candidate.optimization_objectives_auto == ["other_improvement"]


def test_candidate_suggestions_do_not_change_review_status():
    candidate = Paper(
        title="LLM agent for penetration testing",
        status=PaperStatus.SEMANTIC_HIGH.value,
    )

    classify_papers([candidate], CONFIG, include_candidates=True)

    assert candidate.attack_methods_auto == ["penetration_testing"]
    assert candidate.status == PaperStatus.SEMANTIC_HIGH.value


def test_prisma_uses_search_events_and_reason_codes():
    excluded = Paper(status=PaperStatus.HUMAN_EXCLUDED.value, sources=["openalex"])
    excluded.add_screening_decision(
        "human_screening:title_abstract", "exclude", "Out of scope", reason_code="EC2"
    )
    included = Paper(status=PaperStatus.INCLUDED.value, sources=["citation"])
    included.add_screening_decision(
        "human_screening:title_abstract", "include", "In scope", reason_code="IC1"
    )
    included.add_screening_decision(
        "human_screening:full_text", "include", "Full text included", reason_code="IC1"
    )
    counts = build_prisma_counts(
        [excluded, included],
        snowballing=[{"found": 3}],
        search_events=[{"source": "openalex", "n_returned": "10", "n_new": "2"}],
    )

    assert counts["identified_records"] == 13
    assert counts["identified_by_source"] == {"openalex": 10}
    assert counts["title_abstract_exclusions_by_reason"] == {"EC2": 1}
    assert counts["studies_included"] == 1
    assert counts["counts_are_consistent"]


def test_prisma_selects_one_manifest_run_and_excludes_pilot(tmp_path):
    config = {
        "protocol_version": "test-v1",
        "search": {"sources": ["semantic_scholar"], "api": {}},
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    digest = config_hash(config)
    manifest = {
        "status": "complete",
        "config_hash": digest,
        "stages_completed": ["collect"],
        "last_successful_search_run_id": "official",
    }
    (output / "run_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    common = {
        "timestamp": "t", "activity": "search", "source": "semantic_scholar",
        "query": "q", "year_filter": "2022-2026", "n_new": 0,
        "cache_used": "false", "protocol_version": "test-v1", "config_hash": digest,
    }
    log = output / "search_log.csv"
    append_search_log(log, dict(common, run_id="trial", n_returned=40))
    append_search_log(log, dict(common, run_id="official", n_returned=100))
    append_search_log(log, dict(common, run_id="pilot", activity="pilot", n_returned=20))

    pipeline = Pipeline(config_path)
    pipeline.run_prisma()
    default_counts = json.loads((output / "prisma_counts.json").read_text(encoding="utf-8"))
    assert default_counts["identified_records"] == 100

    pipeline.run_prisma(run_id="trial")
    explicit_counts = json.loads((output / "prisma_counts.json").read_text(encoding="utf-8"))
    assert explicit_counts["identified_records"] == 40


def test_prisma_selection_rejects_missing_or_invalid_run():
    events = [
        {"run_id": "bad", "activity": "search", "source": "openalex", "n_returned": "5", "error": "HTTP 429"}
    ]
    for run_id in (None, "missing", "bad"):
        try:
            select_search_events(events, run_id)
        except ValueError:
            pass
        else:
            raise AssertionError(f"run id {run_id!r} should be rejected")


def test_prepare_review_skips_prisma_without_successful_search_run(tmp_path, capsys):
    config = {
        "protocol_version": "test-v1",
        "search": {"sources": ["openalex"], "api": {}},
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    pipeline = Pipeline(config_path)
    completed = []
    pipeline.run_normalize = lambda: completed.append("normalize")
    pipeline.run_dedup = lambda: completed.append("dedup")
    pipeline.run_rule_filter = lambda: completed.append("rule_filter")
    pipeline.run_semantic_filter = lambda: completed.append("semantic_filter")
    pipeline.run_classification = lambda include_candidates=False: completed.append("classification")
    pipeline.run_screening_export = lambda stage: completed.append("screening_export")
    pipeline.run_export = lambda: completed.append("export")

    def unexpected_prisma():
        raise AssertionError("PRISMA must not run without a successful search")

    pipeline.run_prisma = unexpected_prisma
    pipeline.run_stage("prepare_review")

    assert completed == [
        "normalize", "dedup", "rule_filter", "semantic_filter",
        "classification", "screening_export", "export",
    ]
    assert "PRISMA SKIPPED" in capsys.readouterr().out
    manifest = json.loads((tmp_path / "output" / "run_manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "complete"
    assert not (tmp_path / "output" / "prisma_counts.json").exists()


def test_pipeline_rewrites_output_guide_for_each_run(tmp_path):
    config = {
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")

    first_pipeline = Pipeline(config_path)
    guide_path = tmp_path / "output" / "OUTPUT_GUIDE.txt"
    first_run_id = first_pipeline.run_id
    first_guide = guide_path.read_text(encoding="utf-8")
    assert first_run_id in first_guide
    assert "screening_sheet.xlsx" in first_guide
    assert "crosstab.xlsx" in first_guide

    guide_path.unlink()
    guide_path.parent.rmdir()
    second_pipeline = Pipeline(config_path)
    second_guide = guide_path.read_text(encoding="utf-8")

    assert second_pipeline.run_id != first_run_id
    assert second_pipeline.run_id in second_guide
    assert guide_path.exists()


def test_config_hash_tracks_token_usage_but_redacts_credentials():
    base = {
        "taxonomy": {"optimization_objectives": {"token_usage": {"keywords": ["token"]}}},
        "search": {"api": {"semantic_scholar_api_key": "key-a", "contact_email": "a@example.org"}},
    }
    changed_taxonomy = {
        "taxonomy": {"optimization_objectives": {"token_usage": {"keywords": ["tokens consumed"]}}},
        "search": {"api": {"semantic_scholar_api_key": "key-a", "contact_email": "a@example.org"}},
    }
    changed_credentials = {
        "taxonomy": {"optimization_objectives": {"token_usage": {"keywords": ["token"]}}},
        "search": {"api": {"semantic_scholar_api_key": "key-b", "contact_email": "b@example.org"}},
    }

    assert config_hash(base) != config_hash(changed_taxonomy)
    assert config_hash(base) == config_hash(changed_credentials)


def test_analysis_uses_manual_codes_only_for_finally_included():
    included = Paper(
        status=PaperStatus.INCLUDED.value,
        attack_methods_manual=["penetration_testing"],
        optimization_techniques_manual=["tool_use"],
        optimization_objectives_manual=["success_rate"],
    )
    semantic_only = Paper(
        status=PaperStatus.SEMANTIC_HIGH.value,
        attack_methods_manual=["ctf_solving"],
        optimization_techniques_manual=["tool_use"],
    )
    report = build_crosstab([included, semantic_only], CONFIG)

    assert report["total_included"] == 1
    assert report["matrix"]["tool_use"]["penetration_testing"]["k"] == 1
    assert report["objective_matrix"]["tool_use"]["success_rate"]["k"] == 1
    assert all(cell["category"] != "ctf_solving" for cell in report["sparse_cells"])


def test_incomplete_coding_uses_coded_n_and_suppresses_sparse_cells(tmp_path):
    coded = [
        Paper(
            status=PaperStatus.INCLUDED.value,
            attack_methods_manual=["penetration_testing"],
            optimization_techniques_manual=["tool_use"],
            optimization_objectives_manual=["success_rate"],
        )
        for _ in range(3)
    ]
    uncoded = [Paper(status=PaperStatus.INCLUDED.value) for _ in range(2)]
    report = build_crosstab(coded + uncoded, CONFIG)

    method_coverage = report["matrix_coverage"]["optimization_technique_x_attack_method"]
    cell = report["matrix"]["tool_use"]["penetration_testing"]
    assert report["total_included"] == 5
    assert method_coverage["coded_n"] == 3 and method_coverage["complete"] is False
    assert cell["total_n"] == 3
    assert report["sparse_cells"] == []

    workbook_path = tmp_path / "crosstab.xlsx"
    export_crosstab_to_excel(report, workbook_path, gaps=report["sparse_cells"])
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    method_sheet = workbook["Technique x Method"]
    assert "3/5 coded" in method_sheet["A1"].value
    assert "CODING INCOMPLETE" in method_sheet["A2"].value
    assert method_sheet["F4"].value == 3
    assert "Coding incomplete" in workbook["Candidate Sparse Cells"]["A2"].value
    workbook.close()


def test_manual_code_import_and_exact_agreement():
    paper = Paper(
        status=PaperStatus.INCLUDED.value,
        attack_methods_auto=["penetration_testing"],
        optimization_techniques_auto=["tool_use"],
        optimization_objectives_auto=["success_rate"],
    )
    count = import_manual_coding(
        [paper],
        [{
            "id": paper.id,
            "attack_methods_manual": "penetration_testing",
            "optimization_techniques_manual": "tool_use",
            "optimization_objectives_manual": "success_rate",
            "baseline_compared": "manual baseline",
            "measured": "yes",
        }],
        CONFIG,
    )
    agreement = auto_manual_agreement([paper])

    assert count == 1
    assert paper.measured == "yes"
    assert agreement["attack_methods"]["exact_agreement_auto_manual"] == 1.0
    assert agreement["attack_methods"]["cohen_kappa_auto_manual"] == 1.0


def test_full_run_stops_after_search_errors_and_marks_manifest_invalid(tmp_path, monkeypatch):
    config = {
        "protocol_version": "test-v1",
        "search": {"sources": [], "queries": [], "api": {}},
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    pipeline = Pipeline(config_path)
    normalized = []

    def failed_search():
        pipeline.run_errors.append({"stage": "search", "error": "fixture failure"})

    monkeypatch.setattr(pipeline, "run_search", failed_search)
    monkeypatch.setattr(pipeline, "run_normalize", lambda: normalized.append(True))

    try:
        pipeline.run_all()
    except RuntimeError:
        pass
    else:
        raise AssertionError("a search error must invalidate the full run")

    manifest = yaml.safe_load((tmp_path / "output" / "run_manifest.json").read_text())
    assert manifest["valid"] is False
    assert not normalized


def test_validation_sheet_lists_sparse_cells_for_human_search(tmp_path):
    crosstab = {
        "matrix": {},
        "objective_matrix": {},
        "sparse_cells": [{
            "dimension": "optimization_technique_x_attack_method",
            "optimization_technique": "tool_use",
            "category": "penetration_testing",
            "k": 1,
            "row_n": 20,
            "column_n": 10,
            "total_n": 25,
            "papers": ["paper-1"],
        }],
    }
    path = tmp_path / "crosstab.xlsx"
    export_crosstab_to_excel(crosstab, path, gaps=crosstab["sparse_cells"])
    workbook = load_workbook(path, data_only=True)

    sheet = workbook["Validation Search"]
    assert sheet["A2"].value == "optimization_technique_x_attack_method"
    assert sheet["B2"].value == "tool_use x penetration_testing"
    assert sheet["C2"].value is None


def test_crossref_contact_email_uses_mailto_query_parameter():
    adapter = CrossrefAdapter({"search": {"api": {"contact_email": "review@example.org"}}})
    with patch.object(adapter, "_make_request", return_value={"message": {"items": []}}) as request:
        adapter.search("LLM penetration testing", max_results=1)

    assert request.call_args.kwargs["params"]["mailto"] == "review@example.org"


def test_semantic_scholar_search_requests_only_parsed_fields():
    adapter = SemanticScholarAdapter({"search": {"api": {}}})
    with patch.object(adapter, "_make_request", return_value={"data": [], "total": 0}) as request:
        adapter.search("LLM penetration testing", max_results=100)

    fields = request.call_args.kwargs["params"]["fields"].split(",")
    assert fields == [
        "title", "authors", "year", "abstract", "externalIds", "url",
        "venue", "publicationTypes",
    ]


def test_search_discards_records_outside_year_range_and_missing_year(tmp_path):
    config = {
        "search": {
            "sources": ["fixture"],
            "queries": ["LLM security"],
            "max_results_per_query": 10,
            "year_from": 2024,
            "year_to": 2026,
        },
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")

    class FixtureAdapter:
        def search(self, query, max_results):
            return [
                Paper(title="Old paper", year=2023),
                Paper(title="Current paper", year=2024),
                Paper(title="Unknown-year paper"),
                Paper(title="Future paper", year=2027),
            ]

    pipeline = Pipeline(config_path)
    with patch("filtertool.search.get_adapter", return_value=FixtureAdapter()):
        pipeline.run_search()

    papers = pipeline.store.get_all()
    assert [paper.title for paper in papers] == ["Current paper"]
    raw_records = [
        json.loads(line)
        for line in (tmp_path / "output" / "raw_search_results.jsonl").read_text(
            encoding="utf-8"
        ).splitlines()
    ]
    assert [record["record"]["title"] for record in raw_records] == ["Current paper"]
    pipeline.store._conn.close()


def test_dedup_summary_counts_existing_canonical_records(tmp_path, capsys):
    config = {
        "dedup": {"title_fuzzy": {"enabled": False}},
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    pipeline = Pipeline(config_path)
    canonical = Paper(title="Canonical paper", status=PaperStatus.SEMANTIC_HIGH.value)
    duplicate = Paper(
        title="Duplicate paper",
        status=PaperStatus.DUPLICATE.value,
        duplicate_of=canonical.id,
    )
    pipeline.store.add_many([canonical, duplicate])
    pipeline.store.save()

    pipeline.run_dedup()

    assert "[DEDUP DONE] 1 unique, 1 duplicates" in capsys.readouterr().out
    pipeline.store._conn.close()


def test_full_reset_preserves_human_screening_audit(tmp_path):
    config = {
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    store = PaperStore(config["storage"]["database_file"], config["storage"]["cache_dir"])
    reviewed = Paper(
        title="Reviewed", status=PaperStatus.HUMAN_INCLUDED.value,
        human_decision="Include", reason_code="IC1",
    )
    reviewed.add_screening_decision(
        "human_screening:title_abstract", "include", "In scope", reason_code="IC1"
    )
    unreviewed = Paper(title="Unreviewed", status=PaperStatus.SEMANTIC_HIGH.value)
    unreviewed.add_screening_decision("semantic_filter", "HIGH", "score")
    store.add_many([reviewed, unreviewed])
    store.save()

    result = CliRunner().invoke(
        cli, ["reset", "--config", str(config_path)], input="y\n"
    )
    assert result.exit_code == 0, result.output
    store = PaperStore(config["storage"]["database_file"], config["storage"]["cache_dir"])

    assert store.get(reviewed.id).status == PaperStatus.HUMAN_INCLUDED.value
    assert store.get(reviewed.id).screening_decisions[-1]["reason_code"] == "IC1"
    assert store.get(unreviewed.id).status == PaperStatus.NEW.value
    assert store.get(unreviewed.id).screening_decisions


def test_verification_retries_respect_retry_after_and_spacing():
    throttled = type(
        "Response",
        (),
        {"status_code": 429, "headers": {"Retry-After": "12"}},
    )()
    success = type(
        "Response",
        (),
        {"status_code": 200, "headers": {}, "json": lambda self: {"ok": True}},
    )()

    with (
        patch("filtertool.verify.requests.get", side_effect=[throttled, success]) as request,
        patch("filtertool.rate_limit.time.sleep") as sleep,
    ):
        result = _make_request("https://example.test", 2, 5, 6, 10)

    assert result == {"ok": True}
    assert request.call_count == 2
    assert [call.args[0] for call in sleep.call_args_list] == [6, 12.0]


def test_citation_expansion_errors_are_recorded_once():
    pipeline = Pipeline.__new__(Pipeline)
    pipeline.run_errors = []
    failure = {
        "seed_id": "seed-1",
        "direction": "references",
        "error": "HTTP 429",
    }

    pipeline._record_citation_errors([failure, failure])

    assert len(pipeline.run_errors) == 1
    assert pipeline.run_errors[0]["error"] == "HTTP 429"


def test_semantic_scholar_key_comes_from_environment():
    adapter = SemanticScholarAdapter({"search": {"api": {"semantic_scholar_api_key": None}}})
    with (
        patch.dict("os.environ", {"SEMANTIC_SCHOLAR_API_KEY": "dummy-test-key"}),
        patch.object(adapter, "_make_request", return_value={"data": [], "total": 0}) as request,
    ):
        adapter.search("LLM penetration testing", max_results=100)

    assert request.call_args.kwargs["headers"]["x-api-key"] == "dummy-test-key"


def test_semantic_scholar_endpoints_share_provider_floor():
    config = {
        "search": {
            "api": {
                "rate_limit_delay_seconds": 0,
                "minimum_interval_by_host_seconds": {"api.semanticscholar.org": 1.0},
            }
        }
    }
    search_limiter = get_rate_limiter(
        "https://api.semanticscholar.org/graph/v1/paper/search", config
    )
    citation_limiter = get_rate_limiter(
        "https://api.semanticscholar.org/graph/v1/paper/DOI:10.1/x/references", config
    )

    assert search_limiter is citation_limiter
    assert search_limiter.minimum_interval == 1.0


def test_adaptive_limiter_starts_fast_and_recovers_gradually():
    limiter = AdaptiveRateLimiter(minimum_interval=0, successes_to_recover=2)
    with patch("filtertool.rate_limit.time.sleep") as sleep:
        limiter.wait()
        assert not sleep.called
        limiter.record_throttle(1)
        limiter.wait()
        assert sleep.call_args.args == (1,)
        limiter.record_success()
        limiter.record_success()

    assert limiter.current_interval == 0.8


def test_search_adapter_adapts_to_retry_after():
    adapter = BaseSearchAdapter({
        "search": {"api": {"retry_attempts": 2, "retry_delay_seconds": 1}}
    })
    throttled = type(
        "Response",
        (),
        {"status_code": 429, "headers": {"Retry-After": "2"}, "raise_for_status": lambda self: None},
    )()
    success = type(
        "Response",
        (),
        {
            "status_code": 200,
            "headers": {"Content-Type": "application/json"},
            "json": lambda self: {"ok": True},
            "raise_for_status": lambda self: None,
        },
    )()

    with (
        patch("filtertool.search.base.requests.get", side_effect=[throttled, success]),
        patch("filtertool.rate_limit.time.sleep") as sleep,
    ):
        result = adapter._make_request("https://adaptive-test.example/api")

    assert result == {"ok": True}
    assert [call.args[0] for call in sleep.call_args_list] == [2.0]


def test_search_adapter_retries_429_beyond_transient_retry_limit():
    adapter = BaseSearchAdapter({
        "search": {"api": {
            "retry_attempts": 2,
            "rate_limit_retry_attempts": 4,
            "retry_delay_seconds": 1,
        }}
    })
    throttled = type(
        "Response",
        (),
        {
            "status_code": 429,
            "headers": {},
            "raise_for_status": lambda self: None,
        },
    )()
    success = type(
        "Response",
        (),
        {
            "status_code": 200,
            "headers": {"Content-Type": "application/json"},
            "json": lambda self: {"ok": True},
            "raise_for_status": lambda self: None,
        },
    )()
    with (
        patch(
            "filtertool.search.base.requests.get",
            side_effect=[throttled, throttled, throttled, success],
        ) as request,
        patch("filtertool.rate_limit.time.sleep"),
    ):
        result = adapter._make_request("https://retry-429-test.example/api")

    assert result == {"ok": True}
    assert request.call_count == 4


def test_run_continues_and_marks_manifest_partial_for_search_warnings(tmp_path, monkeypatch):
    config = {
        "protocol_version": "test-v1",
        "search": {"sources": [], "queries": [], "api": {}},
        "storage": {
            "database_file": str(tmp_path / "papers.sqlite"),
            "cache_dir": str(tmp_path / "cache"),
        },
        "export": {"output_dir": str(tmp_path / "output")},
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    pipeline = Pipeline(config_path)
    completed = []

    def partial_search():
        pipeline.run_warnings.append({
            "stage": "search",
            "source": "semantic_scholar",
            "query": "fixture",
            "error": "HTTPError: 429",
            "run_id": pipeline.run_id,
        })
        pipeline.last_successful_search_run_id = pipeline.run_id

    monkeypatch.setattr(pipeline, "run_search", partial_search)
    monkeypatch.setattr(pipeline, "run_normalize", lambda: completed.append("normalize"))
    monkeypatch.setattr(pipeline, "run_dedup", lambda: completed.append("dedup"))
    monkeypatch.setattr(pipeline, "run_rule_filter", lambda: completed.append("rule_filter"))
    monkeypatch.setattr(pipeline, "run_semantic_filter", lambda: completed.append("semantic_filter"))
    monkeypatch.setattr(
        pipeline, "run_classification",
        lambda include_candidates=False: completed.append(
            f"classification:{include_candidates}"
        ),
    )
    monkeypatch.setattr(
        pipeline, "run_screening_export",
        lambda stage="title_abstract": completed.append("screening_export"),
    )
    monkeypatch.setattr(pipeline, "run_prisma", lambda run_id=None: completed.append("prisma"))
    monkeypatch.setattr(pipeline, "run_export", lambda: completed.append("export"))

    pipeline.run_all()

    manifest = json.loads(
        (tmp_path / "output" / "run_manifest.json").read_text(encoding="utf-8")
    )
    assert completed == [
        "normalize", "dedup", "rule_filter", "semantic_filter", "classification:True",
        "screening_export", "prisma", "export",
    ]
    assert manifest["status"] == "partial"
    assert manifest["valid"] is True
    assert manifest["errors"] == []
    assert len(manifest["warnings"]) == 1
    pipeline.store._conn.close()


def test_prisma_can_use_successful_events_from_partial_search_run():
    events = [
        {"run_id": "partial", "activity": "search", "n_returned": 0, "error": "HTTP 429"},
        {"run_id": "partial", "activity": "search", "n_returned": 12, "error": ""},
        {"run_id": "other", "activity": "search", "n_returned": 20, "error": ""},
    ]

    selected = select_search_events(events, "partial", allow_partial=True)

    assert selected == [events[1]]
    try:
        select_search_events(events, "partial")
    except ValueError:
        pass
    else:
        raise AssertionError("partial search events must require explicit opt-in")
