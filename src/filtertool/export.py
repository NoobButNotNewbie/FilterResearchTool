"""Excel and JSON export for FilterTool."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from filtertool.models import Paper


# ---------------------------------------------------------------------------
# Color scheme for semantic labels
# ---------------------------------------------------------------------------
_FILLS = {
    "semantic_high": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),   # green
    "semantic_review": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),  # yellow
    "semantic_low": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),     # red
    "verified": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "unverified": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "verify_error": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "duplicate": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),       # grey
    "rule_excluded": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}

_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


# ---------------------------------------------------------------------------
# Column definitions
# ---------------------------------------------------------------------------

# Map of column key -> (header label, width, extractor function)
_COLUMN_DEFS: dict[str, tuple[str, int, Any]] = {
    "id":                     ("ID", 38, lambda p: p.id),
    "title":                  ("Title", 50, lambda p: p.title),
    "authors":                ("Authors", 30, lambda p: "; ".join(p.authors[:5]) + ("..." if len(p.authors) > 5 else "")),
    "year":                   ("Year", 8, lambda p: p.year),
    "doi":                    ("DOI", 25, lambda p: p.doi or ""),
    "url":                    ("URL", 40, lambda p: p.url or ""),
    "sources":                ("Sources", 20, lambda p: ", ".join(p.sources)),
    "status":                 ("Status", 16, lambda p: p.status),
    "semantic_score":         ("Semantic Score", 14, lambda p: round(p.semantic_score, 4) if p.semantic_score is not None else ""),
    "semantic_label":         ("Semantic Label", 14, lambda p: p.semantic_label or ""),
    "keyword_hits":           ("Keyword Hits", 30, lambda p: ", ".join(p.keyword_hits[:10])),
    "keyword_hit_count":      ("Keyword Count", 13, lambda p: p.keyword_hit_count),
    "attack_methods":         ("Attack Methods", 25, lambda p: ", ".join(p.attack_methods)),
    "optimization_techniques": ("Optimization Techniques", 30, lambda p: ", ".join(p.optimization_techniques)),
    "attack_stages":          ("Attack Stages", 25, lambda p: ", ".join(p.attack_stages)),
    "screening_decisions":    ("Screening Log", 50, lambda p: _format_screening(p.screening_decisions)),
    "verification_results":   ("Verification", 40, lambda p: _format_verification(p.verification_results)),
    "venue":                  ("Venue", 30, lambda p: p.venue or ""),
    "abstract":               ("Abstract", 60, lambda p: (p.abstract or "")[:500]),
    "publication_type":        ("Publication Type", 28, lambda p: p.publication_type or ""),
    "human_decision":          ("Human Decision", 18, lambda p: p.human_decision or ""),
    "reason_code":             ("Reason Code", 14, lambda p: p.reason_code or ""),
    "baseline_compared":       ("Baseline Compared", 34, lambda p: p.baseline_compared or ""),
    "measured":                ("Improvement Measured", 20, lambda p: p.measured or ""),
    "attack_methods_auto":      ("Attack Methods (Auto)", 28, lambda p: "; ".join(p.attack_methods_auto)),
    "attack_methods_manual":   ("Attack Methods (Manual)", 30, lambda p: "; ".join(p.attack_methods_manual)),
    "attack_stages_auto":      ("Attack Stages (Auto)", 26, lambda p: "; ".join(p.attack_stages_auto)),
    "attack_stages_manual":    ("Attack Stages (Manual)", 28, lambda p: "; ".join(p.attack_stages_manual)),
    "optimization_techniques_auto": ("Optimization Techniques (Auto)", 36, lambda p: "; ".join(p.optimization_techniques_auto)),
    "optimization_techniques_manual": ("Optimization Techniques (Manual)", 38, lambda p: "; ".join(p.optimization_techniques_manual)),
    "optimization_objectives_auto": ("Optimization Objectives (Auto)", 36, lambda p: "; ".join(p.optimization_objectives_auto)),
    "optimization_objectives_manual": ("Optimization Objectives (Manual)", 38, lambda p: "; ".join(p.optimization_objectives_manual)),
}


def _format_screening(decisions: list[dict]) -> str:
    """Format screening decisions as compact string."""
    parts = []
    for d in decisions:
        s = f"[{d.get('stage', '?')}] {d.get('decision', '?')}"
        reason = d.get("reason", "")
        if reason:
            s += f": {reason[:80]}"
        parts.append(s)
    return " | ".join(parts)


def _format_verification(results: dict[str, dict]) -> str:
    """Format verification results as compact string."""
    parts = []
    for source, r in results.items():
        verified = r.get("verified", False)
        err = r.get("error")
        if err:
            parts.append(f"{source}: ERROR")
        elif verified:
            parts.append(f"{source}: OK")
        else:
            matched = r.get("matched_fields", [])
            mismatched = r.get("mismatched_fields", [])
            parts.append(f"{source}: FAIL({','.join(mismatched)})")
    return " | ".join(parts)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_to_excel(
    papers: list[Paper],
    output_path: str | Path,
    columns: list[str] | None = None,
    sheet_name: str = "Papers",
) -> Path:
    """Export papers to Excel with formatting, URLs, and color coding.
    
    Args:
        papers: List of papers to export
        output_path: Where to save the .xlsx file
        columns: List of column keys to include (default: all)
        sheet_name: Name of the worksheet
    
    Returns:
        Path to the created file
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if columns is None:
        columns = list(_COLUMN_DEFS.keys())

    # Filter to valid columns
    columns = [c for c in columns if c in _COLUMN_DEFS]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # --- Header row ---
    for col_idx, col_key in enumerate(columns, 1):
        header_label, width, _ = _COLUMN_DEFS[col_key]
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # --- Data rows ---
    for row_idx, paper in enumerate(papers, 2):
        for col_idx, col_key in enumerate(columns, 1):
            _, _, extractor = _COLUMN_DEFS[col_key]
            value = extractor(paper)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(col_key in ("title", "abstract", "screening_decisions")))

            # URL column → make it a hyperlink
            if col_key == "url" and value:
                cell.hyperlink = str(value)
                cell.font = Font(color="0563C1", underline="single")

            # DOI column → hyperlink to doi.org
            if col_key == "doi" and value:
                cell.hyperlink = f"https://doi.org/{value}"
                cell.font = Font(color="0563C1", underline="single")

        # Row color based on status
        fill = _FILLS.get(paper.status)
        if fill:
            for col_idx in range(1, len(columns) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    if papers:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(papers) + 1}"

    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Cross-tab Excel export
# ---------------------------------------------------------------------------

def export_crosstab_to_excel(
    crosstab: dict,
    output_path: str | Path,
    sheet_name: str = "CrossTab",
    gaps: list[dict] | None = None,
) -> Path:
    """Export manual-code mapping tables and candidate sparse-cell checks."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    coding_coverage = crosstab.get("coding_coverage", {})
    coverage_summary = "; ".join(
        f"{dimension}: {data.get('coded_n', 0)}/{data.get('total_included', 0)} coded"
        for dimension, data in coding_coverage.items()
    ) or "Coding coverage unavailable"

    def add_table(title: str, matrix: dict, category_header: str, dimension: str) -> None:
        ws = wb.create_sheet(title)
        headers = ["Optimization Technique", category_header, "k", "Row N", "Column N", "Total N", "k/row N", "k/column N", "k/total N", "Paper IDs"]
        widths = [30, 32, 8, 10, 12, 10, 13, 15, 13, 58]
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(1, 1, f"Coding coverage: {coverage_summary}")
        ws.cell(1, 1).font = Font(bold=True)
        table_coverage = crosstab.get("matrix_coverage", {}).get(dimension, {})
        matrix_complete = table_coverage.get("complete", False)
        status_text = (
            f"Table coded N: {table_coverage.get('coded_n', 0)}/"
            f"{table_coverage.get('total_included', 0)} included"
        )
        if not matrix_complete:
            status_text += " | CODING INCOMPLETE - sparse cells suppressed"
        else:
            status_text += " | coding complete"
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(2, 1, status_text)
        ws.cell(2, 1).font = Font(bold=not matrix_complete)
        for col, (header, width) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(3, col, header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.border = _THIN_BORDER
            ws.column_dimensions[get_column_letter(col)].width = width
        row = 4
        sparse_keys = {
            (item.get("optimization_technique"), item.get("category"))
            for item in crosstab.get("sparse_cells", [])
            if item.get("dimension") == dimension
        }
        for technique, categories in matrix.items():
            for category, data in categories.items():
                values = [
                    technique, category, data["k"], data["row_n"], data["column_n"], data["total_n"],
                    data["row_ratio"], data["column_ratio"], data["total_ratio"], ", ".join(data["papers"]),
                ]
                for col, value in enumerate(values, 1):
                    cell = ws.cell(row, col, value)
                    cell.border = _THIN_BORDER
                    cell.alignment = Alignment(vertical="top", wrap_text=(col in {1, 2, 10}))
                    if col in {7, 8, 9}:
                        cell.number_format = "0.00%"
                if (technique, category) in sparse_keys:
                    for cell in ws[row]:
                        cell.fill = PatternFill("solid", fgColor="FCE4D6")
                row += 1
        ws.freeze_panes = "A4"
        if row > 4:
            ws.auto_filter.ref = f"A3:J{row - 1}"

    add_table(
        "Technique x Method", crosstab.get("matrix", {}), "Attack Method",
        "optimization_technique_x_attack_method",
    )
    add_table(
        "Technique x Objective", crosstab.get("objective_matrix", {}), "Optimization Objective",
        "optimization_technique_x_objective",
    )

    gaps_ws = wb.create_sheet("Candidate Sparse Cells")
    gap_headers = ["Dimension", "Optimization Technique", "Category", "k", "Row N", "Column N", "Total N", "Paper IDs"]
    for col, header in enumerate(gap_headers, 1):
        cell = gaps_ws.cell(1, col, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
    if not crosstab.get("coding_complete", False) and not (gaps or []):
        gaps_ws.merge_cells("A2:H2")
        gaps_ws.cell(2, 1, "Coding incomplete; candidate sparse cells are suppressed until related dimensions are fully coded.")
    gap_start_row = 3 if not crosstab.get("coding_complete", False) and not (gaps or []) else 2
    for row, gap in enumerate(gaps or [], gap_start_row):
        values = [gap.get("dimension"), gap.get("optimization_technique"), gap.get("category"), gap.get("k"), gap.get("row_n"), gap.get("column_n"), gap.get("total_n"), ", ".join(gap.get("papers", []))]
        for col, value in enumerate(values, 1):
            gaps_ws.cell(row, col, value).border = _THIN_BORDER
    gaps_ws.freeze_panes = f"A{gap_start_row}"
    gaps_ws.auto_filter.ref = f"A1:H{max(1, len(gaps or []) + gap_start_row - 1)}"
    for col, width in enumerate([42, 30, 34, 8, 10, 12, 10, 58], 1):
        gaps_ws.column_dimensions[get_column_letter(col)].width = width

    validation_ws = wb.create_sheet("Validation Search")
    validation_headers = ["Dimension", "Cell", "Validation Search Query", "Date", "Results Found", "Conclusion", "Reviewer", "Notes"]
    for col, header in enumerate(validation_headers, 1):
        cell = validation_ws.cell(1, col, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
    for row, gap in enumerate(gaps or [], 2):
        validation_ws.cell(row, 1, gap.get("dimension", ""))
        validation_ws.cell(
            row, 2,
            f"{gap.get('optimization_technique', '')} x {gap.get('category', '')}",
        )
    conclusion_validation = DataValidation(
        type="list", formula1='"confirmed gap,keyword artifact,not meaningful"'
    )
    validation_ws.add_data_validation(conclusion_validation)
    conclusion_validation.add(f"F2:F{max(2, len(gaps or []) + 1)}")
    validation_ws.freeze_panes = "A2"
    validation_ws.auto_filter.ref = f"A1:H{max(1, len(gaps or []) + 1)}"
    for col, width in enumerate([38, 38, 60, 16, 16, 32, 24, 60], 1):
        validation_ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# JSON export
# ---------------------------------------------------------------------------

def export_to_json(papers: list[Paper], output_path: str | Path) -> Path:
    """Export papers to JSON."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    data = [p.to_dict() for p in papers]
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return output_path


