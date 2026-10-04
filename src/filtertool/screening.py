"""Human screening sheet export and decision import."""

from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from filtertool.models import Paper, PaperStatus

SCREENING_COLUMNS = [
    "id", "title", "year", "doi", "url", "abstract", "sources",
    "publication_type", "semantic_label", "semantic_score", "keyword_hits",
    "screening_stage", "human_decision", "reason_code", "note", "reviewer",
    "date", "fulltext_retrieved",
]

_DECISION_LABELS = {
    "include": "Include",
    "exclude": "Exclude",
    "unsure": "Unsure",
}
_REASON_CODES = {
    "include": ("IC1", "IC2"),
    "exclude": ("EC1", "EC2", "EC3"),
}
_REASON_CODE_MEANINGS = {
    "IC1": "LLM/agent directly performs an offensive-security task",
    "IC2": "LLM/agent assists, enables, or optimizes an offensive-security task",
    "EC1": "No LLM or AI-agent role",
    "EC2": "No offensive-security task, or defensive-only use",
    "EC3": "Other out-of-scope reason; explain briefly in note",
}


def _status_value(status: str | PaperStatus) -> str:
    return status.value if isinstance(status, PaperStatus) else str(status)


def _latest_screening_decision(paper: Paper, stage: str) -> dict[str, Any] | None:
    audit_stage = f"human_screening:{stage}"
    for decision in reversed(paper.screening_decisions):
        if decision.get("stage") == audit_stage:
            return decision
    return None


def _publication_type(paper: Paper) -> str:
    if paper.publication_type:
        return paper.publication_type
    if "arxiv" in paper.sources:
        return "preprint (arXiv; peer-review status unknown)"
    return "unknown; verify venue/version"


def export_screening_sheet(
    papers: list[Paper], output_path: str | Path, stage: str = "title_abstract"
) -> Path:
    """Create a human-editable workbook for one screening stage."""
    if stage not in {"title_abstract", "full_text"}:
        raise ValueError("stage must be 'title_abstract' or 'full_text'")

    excluded = {
        PaperStatus.DUPLICATE.value,
        PaperStatus.HUMAN_INCLUDED.value,
        PaperStatus.HUMAN_EXCLUDED.value,
        PaperStatus.FULLTEXT_NOT_RETRIEVED.value,
        PaperStatus.FULLTEXT_EXCLUDED.value,
        PaperStatus.INCLUDED.value,
    }
    if stage == "full_text":
        candidates = [
            paper for paper in papers
            if _status_value(paper.status) in {
                PaperStatus.HUMAN_INCLUDED.value,
                PaperStatus.FULLTEXT_ASSESSED.value,
            }
        ]
    else:
        candidates = [
            paper for paper in papers
            if _status_value(paper.status) not in excluded
            and (
                (latest := _latest_screening_decision(paper, "title_abstract")) is None
                or latest.get("decision") == "unsure"
            )
        ]

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Screening"
    header_fill = PatternFill("solid", fgColor="244A63")
    for column, name in enumerate(SCREENING_COLUMNS, 1):
        cell = sheet.cell(1, column, name)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(len(name) + 2, 16), 52
        )

    for row_number, paper in enumerate(candidates, 2):
        values = [
            paper.id,
            paper.title,
            paper.year,
            paper.doi or "",
            paper.url or "",
            paper.abstract or "",
            "; ".join(paper.sources),
            _publication_type(paper),
            paper.semantic_label or "",
            paper.semantic_score,
            "; ".join(paper.keyword_hits),
            stage,
            "",
            "",
            "",
            "",
            "",
            "",
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)

    decision_validation = DataValidation(
        type="list", formula1='"Include,Exclude,Unsure"',
        allow_blank=True, errorStyle="stop", showErrorMessage=True,
    )
    sheet.add_data_validation(decision_validation)
    decision_validation.add(f"M2:M{max(2, len(candidates) + 1)}")

    codebook = workbook.create_sheet("Codebook")
    codebook.append(["Include codes", "Exclude codes", "Unsure"])
    for row_number in range(max(len(_REASON_CODES["include"]), len(_REASON_CODES["exclude"]))):
        codebook.append([
            _REASON_CODES["include"][row_number]
            if row_number < len(_REASON_CODES["include"]) else None,
            _REASON_CODES["exclude"][row_number]
            if row_number < len(_REASON_CODES["exclude"]) else None,
            None,
        ])
    codebook.append([])
    codebook.append(["Decision", "Reason code", "Meaning"])
    for decision, codes in (("Include", _REASON_CODES["include"]), ("Exclude", _REASON_CODES["exclude"])):
        for code in codes:
            codebook.append([decision, code, _REASON_CODE_MEANINGS[code]])
    for cell in codebook[1] + codebook[6]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    codebook.column_dimensions["A"].width = 20
    codebook.column_dimensions["B"].width = 20
    codebook.column_dimensions["C"].width = 72
    workbook.defined_names.add(DefinedName(
        "IncludeReasonCodes", attr_text="'Codebook'!$A$2:$A$3"
    ))
    workbook.defined_names.add(DefinedName(
        "ExcludeReasonCodes", attr_text="'Codebook'!$B$2:$B$4"
    ))
    workbook.defined_names.add(DefinedName(
        "UnsureReasonCodes", attr_text="'Codebook'!$C$2:$C$2"
    ))
    reason_validation = DataValidation(
        type="list",
        formula1='=INDIRECT(IF($M2="Include","IncludeReasonCodes",IF($M2="Exclude","ExcludeReasonCodes","UnsureReasonCodes")))',
        allow_blank=True, errorStyle="stop", showErrorMessage=True,
        errorTitle="Invalid reason code",
        error="Choose a reason code from the list for the selected decision.",
    )
    sheet.add_data_validation(reason_validation)
    reason_validation.add(f"N2:N{max(2, len(candidates) + 1)}")

    if stage == "full_text":
        retrieved_validation = DataValidation(
            type="list", formula1='"yes,no"',
            allow_blank=True, errorStyle="stop", showErrorMessage=True,
        )
        sheet.add_data_validation(retrieved_validation)
        retrieved_validation.add(f"R2:R{max(2, len(candidates) + 1)}")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(1, len(candidates) + 1)}"
    for row in sheet.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output_path)
    return output_path


def read_screening_rows(path: str | Path) -> list[dict[str, Any]]:
    """Read screening decisions from a CSV or the workbook screening sheet."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return [_normalize_row(row) for row in csv.DictReader(stream)]
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("screening input must be .xlsx, .xlsm, or .csv")

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Screening"] if "Screening" in workbook.sheetnames else workbook.active
    values = sheet.iter_rows(values_only=True)
    headers = next(values, None)
    if not headers:
        return []
    rows = []
    for values_row in values:
        if not any(value is not None and str(value).strip() for value in values_row):
            continue
        rows.append(_normalize_row(dict(zip(headers, values_row))))
    return rows


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {}
    for key, value in row.items():
        if key is None:
            continue
        key = re.sub(r"[^a-z0-9_]+", "_", str(key).strip().lower()).strip("_")
        normalized[key] = value.strip() if isinstance(value, str) else value
    return normalized


def apply_screening_rows(papers: list[Paper], rows: list[dict[str, Any]]) -> int:
    """Validate all imported decisions before mutating paper records."""
    by_id = {paper.id: paper for paper in papers}
    prepared = []
    seen_ids = set()

    for row_number, row in enumerate(rows, 2):
        paper_id = str(row.get("id") or "").strip()
        decision = str(row.get("human_decision") or "").strip().lower()
        if not decision:
            continue
        if not paper_id or paper_id not in by_id:
            raise ValueError(f"row {row_number}: unknown or empty paper id")
        if paper_id in seen_ids:
            raise ValueError(f"row {row_number}: duplicate paper id {paper_id}")
        seen_ids.add(paper_id)

        stage = str(row.get("screening_stage") or "title_abstract").strip().lower()
        if stage not in {"title_abstract", "full_text"}:
            raise ValueError(f"row {row_number}: invalid screening_stage {stage!r}")
        if decision not in _DECISION_LABELS:
            raise ValueError(f"row {row_number}: human_decision must be Include, Exclude, or Unsure")

        paper = by_id[paper_id]
        status = _status_value(paper.status)
        if stage == "full_text" and status not in {
            PaperStatus.HUMAN_INCLUDED.value,
            PaperStatus.FULLTEXT_ASSESSED.value,
        }:
            raise ValueError(f"row {row_number}: full-text decision requires HUMAN_INCLUDED")

        reason_code = str(row.get("reason_code") or "").strip().upper()
        if decision in {"include", "exclude"} and reason_code not in _REASON_CODES[decision]:
            options = ", ".join(_REASON_CODES[decision])
            raise ValueError(
                f"row {row_number}: {decision} reason_code must be one of: {options}"
            )

        retrieved = str(row.get("fulltext_retrieved") or "").strip().lower()
        if stage == "full_text" and retrieved not in {"yes", "no"}:
            raise ValueError(f"row {row_number}: fulltext_retrieved must be yes or no")
        prepared.append((paper, row, stage, decision, reason_code, retrieved))

    for paper, row, stage, decision, reason_code, retrieved in prepared:
        human_decision = (
            "Not Retrieved" if stage == "full_text" and retrieved == "no"
            else _DECISION_LABELS[decision]
        )
        decision_or_reason_changed = (
            paper.human_decision != human_decision
            or (paper.reason_code or "") != reason_code
        )
        if stage == "title_abstract":
            status_by_decision = {
                "include": PaperStatus.HUMAN_INCLUDED,
                "exclude": PaperStatus.HUMAN_EXCLUDED,
                "unsure": PaperStatus.REVIEW_NEEDED,
            }
            new_status = status_by_decision[decision]
        elif retrieved == "no":
            new_status = PaperStatus.FULLTEXT_NOT_RETRIEVED
        else:
            status_by_decision = {
                "include": PaperStatus.INCLUDED,
                "exclude": PaperStatus.FULLTEXT_EXCLUDED,
                "unsure": PaperStatus.FULLTEXT_ASSESSED,
            }
            new_status = status_by_decision[decision]

        paper.status = new_status.value
        paper.screening_stage = stage
        paper.human_decision = human_decision
        paper.reason_code = reason_code or None
        paper.screening_note = str(row.get("note") or "").strip() or None
        paper.reviewer = str(row.get("reviewer") or "").strip() or None
        if decision_or_reason_changed or not paper.review_date:
            paper.review_date = datetime.now().astimezone().date().isoformat()
        decision_value = "not_retrieved" if stage == "full_text" and retrieved == "no" else decision
        details = reason_code
        note = paper.screening_note
        reason = "; ".join(value for value in (details, note) if value) or decision_value
        paper.add_screening_decision(
            stage=f"human_screening:{stage}",
            decision=decision_value,
            reason=reason,
            reason_code=reason_code or None,
        )

    return len(prepared)
